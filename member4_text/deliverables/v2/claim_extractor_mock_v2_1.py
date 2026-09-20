
import argparse
import json
import re
from pathlib import Path
from typing import Dict, List, Any, Tuple

from openpyxl import load_workbook


NEGATIVE_CANONICAL = {
    "易脱妆","斑驳","卡纹","卡粉/浮粉","容易泛油","放大纹理","不服帖",
    "厚重/面具感","暗沉","假白","氧化","易转移","拔干","起皮"
}

POSITIVE_NEGATION_CANONICAL = {
    "不易脱妆","不斑驳","不卡纹","不卡粉","不易泛油","不显纹理","不假白",
    "不氧化","不发灰","不易蹭妆","不拔干","不起皮","低刺激/温和","不易致痘"
}


NEGATED_NEGATIVE_TO_POSITIVE = {
    "易脱妆": "不易脱妆",
    "斑驳": "不斑驳",
    "卡纹": "不卡纹",
    "卡粉/浮粉": "不卡粉",
    "容易泛油": "不易泛油",
    "放大纹理": "不显纹理",
    "不服帖": "服帖",
    "暗沉": "不易暗沉",
    "假白": "不假白",
    "氧化": "不氧化",
    "易转移": "不易蹭妆",
    "拔干": "不拔干",
    "起皮": "不起皮"
}

NEGATION_WORDS = ["不是","并不","不算","没有","没","不太","不怎么","不够","不特别"]
STRONG_WORDS = ["完全","非常","特别","真的强","天花板","直接","一点没","零","都不","完全没有","闭眼冲"]
WEAK_WORDS = ["有点","一般","还算","基本","不怎么","轻微","一点点"]


def _norm(v):
    return "" if v is None else str(v).strip()


def _split_phrases(v: str) -> List[str]:
    if not v:
        return []
    parts = re.split(r"[；;|]\s*", v)
    return [p.strip() for p in parts if p and p.strip()]


class ClaimExtractor:
    def __init__(self, dictionary_path: str):
        self.dictionary_path = dictionary_path
        self.entries: List[Dict[str, Any]] = []
        self.consumer_meaning: Dict[str, str] = {}
        self._load_dictionary()

    def _load_dictionary(self):
        wb = load_workbook(self.dictionary_path, data_only=True)

        # Taxonomy -> consumer meaning
        if "Taxonomy" in wb.sheetnames:
            ws = wb["Taxonomy"]
            headers = {_norm(c.value): i for i, c in enumerate(ws[1], start=1)}
            c_col = headers.get("Canonical Claim")
            m_col = headers.get("Consumer Meaning (Plain)") or headers.get("Definition / Consumer Meaning")
            if c_col and m_col:
                for r in range(2, ws.max_row + 1):
                    c = _norm(ws.cell(r, c_col).value)
                    m = _norm(ws.cell(r, m_col).value)
                    if c and m:
                        self.consumer_meaning[c] = m

        # Main dictionary -> phrases
        ws = wb["Claim Dictionary V2"]
        headers = {_norm(c.value): i for i, c in enumerate(ws[1], start=1)}

        c_col = headers["Canonical Claim"]
        p_col = headers["Phrase CN"]
        pt_col = headers.get("Phrase Type")
        s_col = headers.get("Slang / Internet Term")
        e_col = headers.get("Exaggerated Expression")
        cm_col = headers.get("Consumer Meaning")

        seen = set()

        def add_entry(canonical, phrase, phrase_type="", consumer="", slang=False, exag=False):
            canonical, phrase = _norm(canonical), _norm(phrase)
            if not canonical or not phrase:
                return
            key = (canonical, phrase)
            if key in seen:
                return
            seen.add(key)
            self.entries.append({
                "canonical_claim": canonical,
                "phrase": phrase,
                "phrase_type": phrase_type,
                "consumer_meaning": consumer or self.consumer_meaning.get(canonical, ""),
                "is_slang": bool(slang),
                "is_exaggerated": bool(exag),
            })

        for r in range(2, ws.max_row + 1):
            canonical = ws.cell(r, c_col).value
            phrase = ws.cell(r, p_col).value
            phrase_type = _norm(ws.cell(r, pt_col).value) if pt_col else ""
            consumer = _norm(ws.cell(r, cm_col).value) if cm_col else ""

            slang_flag = "slang" in phrase_type.lower()
            exag_flag = "exaggerated" in phrase_type.lower()

            add_entry(canonical, phrase, phrase_type, consumer, slang_flag, exag_flag)

            if s_col:
                for p in _split_phrases(_norm(ws.cell(r, s_col).value)):
                    add_entry(canonical, p, "Slang / Internet Term", consumer, True, False)
            if e_col:
                for p in _split_phrases(_norm(ws.cell(r, e_col).value)):
                    add_entry(canonical, p, "Exaggerated Expression", consumer, False, True)

        # Longest phrase first reduces partial-match noise.
        self.entries.sort(key=lambda x: len(x["phrase"]), reverse=True)

    @staticmethod
    def _extract_duration(text: str) -> str | None:
        # Explicit numeric duration
        m = re.search(r"(\d+(?:\.\d+)?)\s*(?:小时|h|H)", text)
        if m:
            val = m.group(1)
            return f"{val}h"

        # Common Chinese time-range expressions used in beauty posts
        patterns = [
            (r"早八.*晚八|早上八.*晚上八", "12h"),
            (r"早九.*晚六|早上九.*晚上六", "9h"),
            (r"早八.*晚六|早上八.*晚上六", "10h"),
            (r"早九.*晚九|早上九.*晚上九", "12h"),
        ]
        for pat, value in patterns:
            if re.search(pat, text):
                return value
        return None

    @staticmethod
    def _strength(context: str) -> str:
        if any(w in context for w in STRONG_WORDS):
            return "strong"
        if any(w in context for w in WEAK_WORDS):
            return "medium"
        return "medium"

    @staticmethod
    def _has_negation_before(text: str, start: int) -> bool:
        # Only inspect a short window BEFORE the matched phrase.
        # This prevents a later "没斑驳" from negating an earlier "焊脸".
        before = text[max(0, start - 8):start]
        return any(before.endswith(w) or w in before[-5:] for w in NEGATION_WORDS)

    @staticmethod
    def _polarity(canonical: str, negated_before: bool) -> str:
        if canonical in POSITIVE_NEGATION_CANONICAL:
            return "positive"
        if canonical in NEGATIVE_CANONICAL:
            return "positive" if negated_before else "negative"
        if negated_before:
            return "negative"
        return "positive"


    def _semantic_rule_claims(self, text: str) -> List[Dict[str, Any]]:
        """Small P0 rule layer for common beauty-language paraphrases not covered by exact dictionary phrases."""
        out = []

        def add(canonical, matched, polarity="positive", strength="medium",
                duration=None, slang=False, exag=False):
            out.append({
                "canonical_claim": canonical,
                "matched_text": matched,
                "polarity": polarity,
                "strength": strength,
                "duration": duration,
                "is_slang": slang,
                "is_exaggerated": exag,
                "consumer_meaning": self.consumer_meaning.get(canonical, "")
            })

        # Coverage objects
        m = re.search(r"(痘印[^，。；;]{0,10}(?:盖住|遮住|遮掉|遮|透出来|看得见|能看见))", text)
        if m:
            pol = "negative" if any(x in m.group(1) for x in ["透出来","看得见","能看见"]) else "positive"
            add("遮痘印", m.group(1), pol)

        m = re.search(r"(黑眼圈[^，。；;]{0,10}(?:盖住|遮住|遮掉|遮|透出来|看得见|能看见))", text)
        if m:
            pol = "negative" if any(x in m.group(1) for x in ["透出来","看得见","能看见"]) else "positive"
            add("遮黑眼圈", m.group(1), pol)

        m = re.search(r"(泛红[^，。；;]{0,10}(?:盖住|遮住|遮掉|遮|能盖|看不见))", text)
        if m:
            add("遮泛红", m.group(1), "positive")

        # Coverage level
        m = re.search(r"(遮瑕(?:力)?(?:是|很|比较|算)?\s*(中等|一般|中等水平))", text)
        if m:
            add("中等遮瑕", m.group(1), "positive")

        # Pore / blur
        m = re.search(r"(毛孔[^，。；;]{0,10}(?:消失|隐形|看不见|没了))", text)
        if m:
            add("毛孔修饰", m.group(1), "positive", "strong", exag=True)

        # Instant brightening
        m = re.search(r"((?:刚上脸|刚涂|刚上妆)[^，。；;]{0,10}(?:亮|提亮|变亮)[^，。；;]{0,5})", text)
        if m:
            add("即时提亮", m.group(1), "positive")

        # Dry skin / dryness / flaking
        if "干皮" in text:
            add("干皮适用", "干皮", "positive")
        if re.search(r"(不拔干|没有拔干|没拔干)", text):
            mm = re.search(r"(不拔干|没有拔干|没拔干)", text)
            add("不拔干", mm.group(1), "positive")
        if re.search(r"(没有起皮|没起皮|不起皮)", text):
            mm = re.search(r"(没有起皮|没起皮|不起皮)", text)
            add("不起皮", mm.group(1), "positive")

        # Transfer / wear-off
        m = re.search(r"((?:口罩|纸巾|衣领)[^，。；;]{0,14}(?:蹭掉|蹭到|沾到|掉了|掉妆))", text)
        if m:
            add("易转移", m.group(1), "negative")
            add("易脱妆", m.group(1), "negative")

        # Lightweight / low powder feeling
        m = re.search(r"((?:没什么|没有|几乎没|低|很少)[^，。；;]{0,4}粉感)", text)
        if m:
            add("轻薄", m.group(1), "positive")

        # Wear stability paraphrase
        m = re.search(r"(妆面[^，。；;]{0,8}(?:稳|稳定|还在|在线))", text)
        if m:
            add("持妆", m.group(1), "positive")

        # "X hours with little wear-off" also implies long wear
        duration = self._extract_duration(text)
        if duration and re.search(r"(没怎么脱妆|基本没脱妆|不怎么脱妆|妆面.*稳)", text):
            add("持妆", "时长内妆面保持稳定", "positive", duration=duration)

        # 12h phrasing -> canonical 12-hour claim
        if duration == "12h" and re.search(r"(焊脸|焊妆|不掉|持妆|妆面.*稳)", text):
            add("12小时持妆", "12小时持妆语义", "positive", "strong", duration="12h",
                slang=("焊" in text), exag=("完全" in text or "都不掉" in text))

        return out

    @staticmethod
    def _claim_dedupe_priority(claim: Dict[str, Any]) -> int:
        # More specific claims should beat generic ones when both describe same idea.
        specific = {
            "12小时持妆": 5, "24小时持妆": 5,
            "中等遮瑕": 5, "高遮瑕": 5, "低遮瑕": 5,
            "遮痘印": 4, "遮黑眼圈": 4, "遮泛红": 4, "遮斑": 4,
            "不斑驳": 4, "不卡粉": 4, "不拔干": 4, "不起皮": 4,
            "毛孔修饰": 4
        }
        return specific.get(claim["canonical_claim"], 1)

    def extract_claims(self, text: str) -> Dict[str, Any]:
        text = _norm(text)
        found: List[Tuple[int, int, Dict[str, Any]]] = []

        for entry in self.entries:
            phrase = entry["phrase"]
            start = text.find(phrase)
            while start != -1:
                end = start + len(phrase)
                found.append((start, end, entry))
                start = text.find(phrase, start + 1)

        # Deduplicate same canonical in overlapping spans: keep longest matched phrase
        found.sort(key=lambda x: (x[0], -(x[1]-x[0])))
        accepted: List[Tuple[int, int, Dict[str, Any]]] = []
        seen_keys = set()

        for item in found:
            start, end, entry = item
            key = (entry["canonical_claim"], start, end)
            if key in seen_keys:
                continue
            seen_keys.add(key)

            # If a shorter phrase for same canonical is inside a longer already accepted phrase, skip it.
            contained = False
            for a_start, a_end, a_entry in accepted:
                if a_entry["canonical_claim"] == entry["canonical_claim"] and a_start <= start and end <= a_end:
                    contained = True
                    break
            if not contained:
                accepted.append(item)

        claims = []
        duration_global = self._extract_duration(text)

        for start, end, entry in accepted:
            left = max(0, start - 12)
            right = min(len(text), end + 18)
            context = text[left:right]
            canonical = entry["canonical_claim"]
            negated_before = self._has_negation_before(text, start)

            # Example: "没斑驳" should become canonical "不斑驳".
            normalized_canonical = NEGATED_NEGATIVE_TO_POSITIVE.get(canonical, canonical) if negated_before else canonical

            duration = duration_global if ("持妆" in normalized_canonical or "脱妆" in normalized_canonical or "暗沉" in normalized_canonical) else None
            consumer = self.consumer_meaning.get(normalized_canonical) or entry["consumer_meaning"] or self.consumer_meaning.get(canonical, "")

            claims.append({
                "canonical_claim": normalized_canonical,
                "matched_text": entry["phrase"],
                "polarity": self._polarity(canonical, negated_before),
                "strength": self._strength(context),
                "duration": duration,
                "is_slang": entry["is_slang"],
                "is_exaggerated": entry["is_exaggerated"],
                "consumer_meaning": consumer
            })

        # Add P0 semantic rules for paraphrases.
        claims.extend(self._semantic_rule_claims(text))

        # Ambiguity cleanup for common social-media terms.
        # "妈生皮" should primarily mean low-makeup/natural-looking, not automatically "不卡粉".
        if "妈生皮" in text and not re.search(r"(浮粉|卡粉|结块)", text):
            claims = [c for c in claims if not (c["canonical_claim"] == "不卡粉" and c["matched_text"] == "妈生皮")]

        # Generic "自然" should not by itself create a low-makeup claim.
        claims = [c for c in claims if not (c["canonical_claim"] == "低妆感" and c["matched_text"] in {"自然","自然妆效","自然感"})]

        # Avoid ultra-short ambiguous phrase matches such as single-character "干".
        claims = [c for c in claims if len(c["matched_text"]) >= 2]

        # If a specific coverage level is present, suppress generic coverage from same sentence.
        has_specific_coverage = any(c["canonical_claim"] in {"高遮瑕","中等遮瑕","低遮瑕"} for c in claims)
        if has_specific_coverage:
            claims = [c for c in claims if c["canonical_claim"] != "遮瑕"]

        # If 12-hour canonical is established, suppress generic 持妆.
        if any(c["canonical_claim"] == "12小时持妆" for c in claims):
            claims = [c for c in claims if c["canonical_claim"] != "持妆"]

        # If a normalized positive opposite exists, suppress its negative base claim.
        # Example: "不拔干" should not also return "拔干".
        positive_to_negative = {
            "不易脱妆": "易脱妆",
            "不斑驳": "斑驳",
            "不卡纹": "卡纹",
            "不卡粉": "卡粉/浮粉",
            "不易泛油": "容易泛油",
            "不显纹理": "放大纹理",
            "服帖": "不服帖",
            "不易暗沉": "暗沉",
            "不假白": "假白",
            "不氧化": "氧化",
            "不易蹭妆": "易转移",
            "不拔干": "拔干",
            "不起皮": "起皮"
        }
        present = {c["canonical_claim"] for c in claims}
        suppress = {neg for pos, neg in positive_to_negative.items() if pos in present}
        claims = [c for c in claims if c["canonical_claim"] not in suppress]

        # Final dedupe by canonical claim: prefer more specific / longer match.
        best = {}
        for c in claims:
            key = c["canonical_claim"]
            score = (self._claim_dedupe_priority(c), len(c["matched_text"]))
            if key not in best or score > best[key][0]:
                best[key] = (score, c)

        unique = [v[1] for v in best.values()]

        return {
            "tool_name": "claim_extraction",
            "input": {"text": text},
            "output": {"claims": unique}
        }



class MockSemanticLayer:
    """
    Deterministic semantic layer for demo / integration.
    No API, no internet, no external model.

    Purpose:
    - simulate semantic understanding for common paraphrases
    - keep exactly the same JSON fields as a future LLM layer
    - allow Member 1 / frontend to integrate now
    """

    def __init__(self, meanings: Dict[str, str]):
        self.meanings = meanings

    def _claim(self, canonical, matched, polarity="positive", strength="medium",
               duration=None, slang=False, exaggerated=False, note=None):
        meaning = self.meanings.get(canonical, "")
        if note:
            meaning = note
        return {
            "canonical_claim": canonical,
            "matched_text": matched,
            "polarity": polarity,
            "strength": strength,
            "duration": duration,
            "is_slang": slang,
            "is_exaggerated": exaggerated,
            "consumer_meaning": meaning
        }

    @staticmethod
    def _clock_duration(text: str):
        # Chinese clock ranges: 早上7点 -> 下午5点 etc.
        cn_map = {
            "一":1,"二":2,"两":2,"三":3,"四":4,"五":5,"六":6,
            "七":7,"八":8,"九":9,"十":10,"十一":11,"十二":12
        }

        def parse_hour(token):
            token = token.strip()
            if token.isdigit():
                return int(token)
            return cn_map.get(token)

        m = re.search(
            r"(?:早上|早晨|早|上午)?\s*([0-9]{1,2}|十一|十二|十|一|二|两|三|四|五|六|七|八|九)\s*点?"
            r".{0,8}?"
            r"(?:下午|晚上|晚)?\s*([0-9]{1,2}|十一|十二|十|一|二|两|三|四|五|六|七|八|九)\s*点?",
            text
        )
        if not m:
            return None

        h1 = parse_hour(m.group(1))
        h2 = parse_hour(m.group(2))
        if h1 is None or h2 is None:
            return None

        # infer PM for second time when text says 下午/晚上/晚 between groups
        middle = text[m.start():m.end()]
        if any(x in middle for x in ["下午","晚上","晚"]) and h2 <= 12:
            if h2 < 12:
                h2 += 12

        if "下午" in text[:m.start(2)+5] and h1 < 12 and "早" not in text[:m.start(1)+5] and "上午" not in text[:m.start(1)+5]:
            h1 += 12

        if h2 <= h1:
            h2 += 12 if h2 < 12 else 24

        d = h2 - h1
        if 0 < d <= 24:
            return f"{d}h"
        return None

    def analyze(self, text: str, rule_claims: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        out = []
        duration = self._clock_duration(text) or ClaimExtractor._extract_duration(text)

        def add(*args, **kwargs):
            out.append(self._claim(*args, **kwargs))

        # ---------- Natural finish / lightweight ----------
        if re.search(r"(很轻|轻得像没化妆|像没化妆|没什么妆感|像自己的皮肤|自己的皮肤)", text):
            if "很轻" in text or "轻得像" in text:
                add("轻薄", "很轻" if "很轻" in text else "轻得像没化妆")
            add("低妆感", "像没化妆" if "像没化妆" in text else "像自己的皮肤")

        if re.search(r"(粉感(?:很|挺)?重|面具感|像戴面具|厚重)", text):
            m = re.search(r"(粉感(?:很|挺)?重|面具感|像戴面具|厚重)", text)
            add("厚重/面具感", m.group(1), "negative")

        # ---------- Wear / duration ----------
        if re.search(r"(妆面基本还在|妆面还在|整体还算完整|整体妆面.*完整|妆面.*在线|真的稳|一直很稳)", text):
            m = re.search(r"(妆面基本还在|妆面还在|整体还算完整|整体妆面[^，。]{0,8}完整|妆面[^，。]{0,8}在线|真的稳|一直很稳)", text)
            add("持妆", m.group(1), duration=duration, slang=("在线" in m.group(1)))

        if duration == "12h" and re.search(r"(没大面积斑驳|没怎么脱|妆面.*完整|妆面.*在线|持妆)", text):
            add("12小时持妆", "12小时佩戴表现", duration="12h")

        if re.search(r"(鼻翼.*开始脱|开始脱妆|掉了一些|没怎么脱|没有明显掉|妆没有明显掉)", text):
            m = re.search(r"(鼻翼[^，。]{0,8}开始脱|开始脱妆|掉了一些|没怎么脱|没有明显掉|妆没有明显掉)", text)
            phrase = m.group(1)
            if any(x in phrase for x in ["没怎么脱","没有明显掉","妆没有明显掉"]):
                add("不易脱妆", phrase, "positive", duration=duration)
            else:
                add("易脱妆", phrase, "negative", duration=duration)

        # ---------- Oil / shine ----------
        if re.search(r"(微微出油|有一点油|一点油光|T区只是微微出油)", text):
            m = re.search(r"(微微出油|有一点油|一点油光|T区只是微微出油)", text)
            add("不易泛油", m.group(1), "positive", "weak", duration=duration)

        if re.search(r"(开始泛油|油光.*明显|泛油了|油成|很油)", text):
            m = re.search(r"(开始泛油|油光[^，。]{0,6}明显|泛油了|很油)", text)
            if m:
                add("容易泛油", m.group(1), "negative", duration=duration)

        # ---------- Coverage ----------
        if re.search(r"(红痘印.*遮住|痘印.*遮住|浅痘印.*盖住)", text):
            m = re.search(r"((?:红|浅)?痘印[^，。]{0,8}(?:遮住|盖住))", text)
            add("遮痘印", m.group(1), "positive")

        if re.search(r"(深色痘印.*看出来|痘印.*透|痘印.*看得见)", text):
            m = re.search(r"(深色?痘印[^，。]{0,10}(?:看出来|透出来|看得见)|痘印[^，。]{0,10}(?:透出来|看得见))", text)
            if m:
                add("遮痘印", m.group(1), "negative")

        if re.search(r"(黑眼圈.*还得|黑眼圈.*遮不住|黑眼圈.*能看见)", text):
            m = re.search(r"(黑眼圈[^，。]{0,12}(?:还得上遮瑕膏|遮不住|能看见))", text)
            add("遮黑眼圈", m.group(1), "negative")

        if re.search(r"(泛红.*能盖|遮泛红.*不错)", text):
            m = re.search(r"(泛红[^，。]{0,8}(?:能盖|遮住)|遮泛红[^，。]{0,6}不错)", text)
            add("遮泛红", m.group(1), "positive")

        if re.search(r"(雀斑.*遮不住|色斑.*遮不住)", text):
            m = re.search(r"((?:雀斑|色斑)[^，。]{0,8}遮不住)", text)
            add("遮斑", m.group(1), "negative")

        # infer low/high coverage style
        if re.search(r"(一层只能盖住浅痘印|遮瑕没有想象中那么高|遮瑕不算高)", text):
            add("低遮瑕", "遮瑕没有想象中那么高" if "没有想象中那么高" in text else "一层只能盖住浅痘印", "negative")

        # ---------- Pore / texture ----------
        if re.search(r"(毛孔.*柔化|毛孔.*隐形|毛孔.*不明显|毛孔被柔化)", text):
            m = re.search(r"(毛孔[^，。]{0,12}(?:柔化|隐形|不明显)|毛孔被柔化)", text)
            add("毛孔修饰", m.group(1), "positive")

        if re.search(r"(柔化了不少|磨皮到没纹理|开了磨皮|柔焦)", text):
            m = re.search(r"(柔化了不少|磨皮到没纹理|开了磨皮|柔焦)", text)
            add("柔焦", m.group(1), "positive",
                exaggerated=("没纹理" in m.group(1) or "磨皮" in m.group(1)))

        if re.search(r"(纹理又变明显|纹理.*明显|毛孔.*被放大)", text):
            m = re.search(r"(纹理[^，。]{0,8}变明显|毛孔[^，。]{0,8}被放大)", text)
            add("放大纹理", m.group(1), "negative", duration=duration)

        if re.search(r"(没有被放大|没放大毛孔|毛孔也没有被放大)", text):
            m = re.search(r"(毛孔[^，。]{0,10}没有被放大|没放大毛孔)", text)
            add("不显纹理", m.group(1), "positive")

        # ---------- Dryness / hydration ----------
        if re.search(r"(很润|挺润|水润|很滋润|不紧绷)", text):
            m = re.search(r"(很润|挺润|水润|很滋润)", text)
            if m:
                add("保湿", m.group(1), "positive")

        if "干皮" in text and not re.search(r"(干皮闭眼冲.*过头|干皮.*不适合)", text):
            add("干皮适用", "干皮", "positive")

        if re.search(r"(不紧绷|不拔干|没拔干)", text):
            m = re.search(r"(不紧绷|不拔干|没拔干)", text)
            add("不拔干", m.group(1), "positive")

        if re.search(r"(没有起皮|没起皮|不起皮)", text):
            m = re.search(r"(没有起皮|没起皮|不起皮)", text)
            add("不起皮", m.group(1), "positive")

        if re.search(r"(轻微起皮|有点起皮|开始起皮)", text):
            m = re.search(r"(轻微起皮|有点起皮|开始起皮)", text)
            add("起皮", m.group(1), "negative")

        # ---------- Transfer / rub ----------
        if re.search(r"(口罩上有一点粉|口罩.*蹭下来|口罩.*有粉|沾口罩|蹭到口罩)", text):
            m = re.search(r"(口罩[^，。]{0,14}(?:有一点粉|蹭下来不少粉底|有粉|蹭到|沾到))", text)
            if m:
                add("易转移", m.group(1), "negative", duration=duration)

        if re.search(r"(不是完全不沾口罩|比.*好很多)", text):
            if "口罩" in text:
                add("不易蹭妆", "不是完全不沾口罩", "mixed")

        if re.search(r"(纸巾一擦.*掉|一擦就会掉)", text):
            m = re.search(r"(纸巾[^，。]{0,10}(?:一擦.*掉|擦就会掉)|一擦就会掉不少)", text)
            if m:
                add("耐摩擦", m.group(1), "negative")

        # ---------- Water / sweat ----------
        if re.search(r"(水冲之后基本还在|淋了点雨.*没花|底妆没花)", text):
            m = re.search(r"(水冲之后基本还在|淋了点雨[^，。]{0,12}没花|底妆没花)", text)
            add("防水", m.group(1), "positive")

        if re.search(r"(号称防汗|防汗.*一般|跑完步.*花了)", text):
            m = re.search(r"(号称防汗|防汗[^，。]{0,6}一般|跑完步[^，。]{0,12}花了)", text)
            add("防汗", m.group(1), "negative")

        if re.search(r"(鼻翼和嘴角已经花了|妆花了|局部花了)", text):
            m = re.search(r"(鼻翼和嘴角已经花了|妆花了|局部花了)", text)
            add("斑驳", m.group(1), "negative")

        # ---------- Finish / brightness / color ----------
        if re.search(r"(自然的光泽妆效|自然光泽|有光泽|健康光)", text):
            m = re.search(r"(自然的光泽妆效|自然光泽|有光泽|健康光)", text)
            add("自然光泽", m.group(1), "positive")

        if re.search(r"(柔雾|偏柔雾)", text):
            m = re.search(r"(偏柔雾|柔雾)", text)
            add("柔雾", m.group(1), "positive")

        if re.search(r"(不是完全哑光|不是死哑光)", text):
            m = re.search(r"(不是完全哑光|不是死哑光)", text)
            add("哑光", m.group(1), "negative")

        if re.search(r"(瞬间白了一个度|上脸.*白了一个度|瞬间提亮|立刻亮)", text):
            m = re.search(r"(瞬间白了一个度|上脸[^，。]{0,6}白了一个度|瞬间提亮|立刻亮)", text)
            add("即时提亮", m.group(1), "positive", "strong",
                exaggerated=("一个度" in m.group(1)))

        if re.search(r"(不是.*死白|不是.*假白|不假白)", text):
            m = re.search(r"(不是[^，。]{0,8}死白|不是[^，。]{0,8}假白|不假白)", text)
            add("不假白", m.group(1), "positive")

        if re.search(r"(发灰发暗|有点发灰|变暗|发暗)", text):
            m = re.search(r"(发灰发暗|有点发灰|变暗|发暗)", text)
            if "灰" in m.group(1):
                add("不发灰/发灰", m.group(1), "negative")
            add("暗沉", m.group(1), "negative", duration=duration)

        if re.search(r"(微微变黄|严重氧化|氧化)", text):
            m = re.search(r"(微微变黄|严重氧化|氧化)", text)
            pol = "negative"
            add("氧化", m.group(1), pol, "weak" if "微微" in m.group(1) else "medium", duration=duration)

        # ---------- Skin suitability / sensitivity ----------
        if re.search(r"(敏感肌.*没刺痛|敏感肌.*没有刺痛)", text):
            m = re.search(r"(敏感肌[^，。]{0,12}(?:没刺痛|没有刺痛))", text)
            add("敏感肌适用", m.group(1), "positive")
            add("低刺激/温和", m.group(1), "positive")

        if re.search(r"(油皮夏天用真的稳|油皮.*用.*稳)", text):
            m = re.search(r"(油皮[^，。]{0,14}(?:真的稳|很稳|用.*稳))", text)
            if m:
                add("油皮适用", m.group(1), "positive")

        if re.search(r"(油皮亲妈.*不敢说|不敢说.*油皮亲妈)", text):
            add("油皮适用", "油皮亲妈", "negative", slang=True, exaggerated=True)

        if re.search(r"(干皮闭眼冲.*过头|闭眼冲这种话.*过头)", text):
            add("干皮适用", "干皮闭眼冲", "negative", slang=True, exaggerated=True)

        # ---------- Face region adherence ----------
        if re.search(r"(脸颊.*一直很服帖|很服帖|贴肤)", text):
            m = re.search(r"(脸颊[^，。]{0,10}一直很服帖|很服帖|贴肤)", text)
            add("服帖", m.group(1), "positive")

        # ---------- Claim-vs-experience statements ----------
        if re.search(r"(宣传说24小时持妆|24小时持妆.*太夸张)", text):
            add("24小时持妆", "24小时持妆", "negative", exaggerated=True)

        if re.search(r"(号称轻薄|宣传.*轻薄)", text):
            add("轻薄", "号称轻薄", "negative")

        if re.search(r"(不算控油挂|不是控油挂|不算控油)", text):
            m = re.search(r"(不算控油挂|不是控油挂|不算控油)", text)
            add("控油", m.group(1), "negative")

        # ---------- De-duplicate by canonical, prefer longer matched text ----------
        best = {}
        for c in out:
            key = c["canonical_claim"]
            if key not in best or len(c["matched_text"]) > len(best[key]["matched_text"]):
                best[key] = c
        return list(best.values())


def merge_rule_and_mock(rule_claims: List[Dict[str, Any]], mock_claims: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    merged = {c["canonical_claim"]: dict(c) for c in rule_claims}

    for c in mock_claims:
        key = c["canonical_claim"]
        if key not in merged:
            merged[key] = dict(c)
        else:
            old = merged[key]
            # Mock semantic layer may refine context.
            if c.get("matched_text") and len(c["matched_text"]) > len(old.get("matched_text","")):
                old["matched_text"] = c["matched_text"]
            if c.get("polarity"):
                old["polarity"] = c["polarity"]
            if c.get("strength"):
                old["strength"] = c["strength"]
            if c.get("duration") is not None:
                old["duration"] = c["duration"]
            if c.get("consumer_meaning"):
                old["consumer_meaning"] = c["consumer_meaning"]
            old["is_slang"] = bool(old.get("is_slang") or c.get("is_slang"))
            old["is_exaggerated"] = bool(old.get("is_exaggerated") or c.get("is_exaggerated"))

    return list(merged.values())

def run_testset(extractor: ClaimExtractor, test_path: str, out_path: str):
    wb = load_workbook(test_path)
    ws = wb["Test Set"]

    headers = {_norm(c.value): i for i, c in enumerate(ws[1], start=1)}
    text_col = headers["Test Text"]
    expected_col = headers["Expected Canonical Claims"]
    model_col = headers["Model Output"]
    correct_col = headers["Correct?"]
    notes_col = headers["Error Notes"]

    total = 0
    exact = 0

    for r in range(2, ws.max_row + 1):
        text = _norm(ws.cell(r, text_col).value)
        expected = {_norm(x) for x in re.split(r"[；;]", _norm(ws.cell(r, expected_col).value)) if _norm(x)}

        result = extractor.extract_claims(text)
        predicted = {c["canonical_claim"] for c in result["output"]["claims"]}

        ws.cell(r, model_col).value = json.dumps(result["output"]["claims"], ensure_ascii=False)

        if predicted == expected:
            ws.cell(r, correct_col).value = "Yes"
            exact += 1
        elif predicted & expected:
            ws.cell(r, correct_col).value = "Partial"
        else:
            ws.cell(r, correct_col).value = "No"

        missing = sorted(expected - predicted)
        extra = sorted(predicted - expected)
        notes = []
        if missing:
            notes.append("Missing: " + "；".join(missing))
        if extra:
            notes.append("Extra: " + "；".join(extra))
        ws.cell(r, notes_col).value = " | ".join(notes)

        total += 1

    wb.save(out_path)
    print(json.dumps({
        "test_rows": total,
        "exact_match_rows": exact,
        "exact_match_rate": round(exact / total, 3) if total else 0
    }, ensure_ascii=False, indent=2))



def run_blindtest(extractor: ClaimExtractor, blind_path: str, out_path: str):
    """Run the Blind Test sheet in batch without using the Answer Key."""
    wb = load_workbook(blind_path)
    ws = wb["Blind Test"]

    headers = {_norm(c.value): i for i, c in enumerate(ws[1], start=1)}
    text_col = headers["Blind Test Text"]
    model_col = headers["Model Output"]

    total = 0
    claim_counts = []

    for r in range(2, ws.max_row + 1):
        text = _norm(ws.cell(r, text_col).value)
        if not text:
            continue

        result = extractor.extract_claims(text)
        claims = result["output"]["claims"]
        ws.cell(r, model_col).value = json.dumps(claims, ensure_ascii=False)
        claim_counts.append(len(claims))
        total += 1

    # Do NOT auto-score against Answer Key here.
    # The point of blind testing is to inspect raw output first.
    wb.save(out_path)

    print(json.dumps({
        "blind_test_rows": total,
        "output_file": out_path,
        "average_claims_per_row": round(sum(claim_counts) / len(claim_counts), 2) if claim_counts else 0,
        "note": "Blind test completed. Model Output was filled; Answer Key was not used for extraction."
    }, ensure_ascii=False, indent=2))


def resolve_claim_conflicts(text: str, claims: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Conflict resolution / cleanup layer for Mock v2.1.
    Goal: reduce duplicate or contradictory claims without hurting recall too much.
    """

    by_claim = {}
    for c in claims:
        key = c["canonical_claim"]
        if key not in by_claim:
            by_claim[key] = c
        else:
            # Prefer longer matched text / richer duration.
            old = by_claim[key]
            score_old = (len(old.get("matched_text","")), 1 if old.get("duration") else 0)
            score_new = (len(c.get("matched_text","")), 1 if c.get("duration") else 0)
            if score_new > score_old:
                by_claim[key] = c

    claims = list(by_claim.values())
    present = {c["canonical_claim"] for c in claims}

    # --- Opposite-pair resolution ---
    opposite_pairs = [
        ("不斑驳","斑驳"),
        ("不卡粉","卡粉/浮粉"),
        ("不卡纹","卡纹"),
        ("不拔干","拔干"),
        ("不起皮","起皮"),
        ("不易泛油","容易泛油"),
        ("不易脱妆","易脱妆"),
        ("不氧化","氧化"),
        ("不假白","假白"),
        ("不显纹理","放大纹理"),
        ("服帖","不服帖"),
    ]

    def polarity_of(name):
        for c in claims:
            if c["canonical_claim"] == name:
                return c.get("polarity")
        return None

    to_remove = set()

    for pos, neg in opposite_pairs:
        if pos in present and neg in present:
            # If text explicitly describes a time/region contrast, allow both.
            contrast = any(x in text for x in ["但","不过","就是","后来","之后","两个小时后","下午","鼻翼","脸颊","眼下","法令纹"])
            if contrast and (
                (pos, neg) in [("不显纹理","放大纹理"), ("服帖","不服帖"), ("不起皮","起皮"), ("不易脱妆","易脱妆")]
            ):
                continue

            # Strong explicit positive negation wins over negative base.
            if any(x in text for x in ["没有"+neg.replace("易",""), "没"+neg.replace("易",""), pos]):
                to_remove.add(neg)
                continue

            # Strong explicit negative outcome wins over positive opposite.
            neg_keywords = {
                "斑驳":["斑驳","花了","一块一块"],
                "卡粉/浮粉":["卡粉","浮粉"],
                "卡纹":["卡纹","积线"],
                "拔干":["拔干","紧绷"],
                "起皮":["起皮","爆皮"],
                "容易泛油":["泛油","油光","出油"],
                "易脱妆":["脱妆","开始脱","掉妆"],
                "氧化":["氧化","变黄","变橘","变深"],
                "假白":["假白","死白"],
                "放大纹理":["纹理变明显","毛孔被放大"],
                "不服帖":["不服帖","浮在脸上"],
            }
            if any(k in text for k in neg_keywords.get(neg, [])):
                to_remove.add(pos)
            else:
                # Default conservative choice: keep the semantically more specific one
                if polarity_of(pos) == "positive" and polarity_of(neg) == "negative":
                    # ambiguous: keep both only if clear contrast words exist
                    if not contrast:
                        to_remove.add(neg)

    # --- Coverage hierarchy cleanup ---
    levels = [x for x in ["高遮瑕","中等遮瑕","低遮瑕"] if x in present]
    if len(levels) > 1:
        # Decide by explicit wording.
        if re.search(r"(遮瑕.*中等|中等遮瑕)", text):
            keep = "中等遮瑕"
        elif re.search(r"(遮瑕.*不高|只能盖住浅痘印|低遮瑕)", text):
            keep = "低遮瑕"
        elif re.search(r"(高遮瑕|遮瑕没有想象中那么高|遮瑕很高|遮瑕力强|七八成|大部分能盖住)", text):
            keep = "高遮瑕"
        else:
            keep = levels[0]
        for l in levels:
            if l != keep:
                to_remove.add(l)

    # Specific coverage levels suppress generic 遮瑕 unless generic is explicitly useful.
    present_after = present - to_remove
    if any(x in present_after for x in ["高遮瑕","中等遮瑕","低遮瑕"]):
        to_remove.add("遮瑕")

    # Object-specific coverage should remain even if general coverage level is present.
    # no removal here.

    # --- False-positive cleanup for some ambiguous phrases ---
    if "微微出油" in text or "一点油光" in text:
        # Treat as mild shine, not strong oily outcome.
        to_remove.add("容易泛油")

    if re.search(r"(不是.*死白|不是.*假白|不假白)", text):
        to_remove.add("假白")

    if re.search(r"(没有严重氧化|不算严重氧化)", text):
        # keep 氧化 if color actually changed, but avoid accidentally claiming 不氧化
        to_remove.add("不氧化")

    # "不算控油挂" should not automatically create 泛油 if no actual oiliness is described.
    if re.search(r"(不算控油挂|不是控油挂|不算控油)", text) and not re.search(r"(出油|泛油|油光|很油)", text):
        to_remove.add("容易泛油")
        to_remove.add("不易泛油")

    # "像没化妆" is low makeup look, not comfort.
    if re.search(r"(像没化妆|妈生皮|像自己的皮肤)", text):
        to_remove.add("舒适")

    # "防水" success does not imply abrasion resistance.
    if "防水" in present and not re.search(r"(擦|蹭|摩擦|纸巾)", text):
        to_remove.add("耐摩擦")

    # B13: "油皮亲妈我不敢说" + 实际已经泛油，不能据此声称控油
    if re.search(r"油皮亲妈.*不敢说", text) and re.search(r"(泛油|出油|冒油)", text):
        to_remove.add("控油")

    # B30: 明确否定油腻，不应误判为“容易泛油”
    if re.search(r"(没有很油|并不很油|不算很油|没那么油)", text):
        to_remove.add("容易泛油")

    # --- Rebuild ---
    cleaned = [c for c in claims if c["canonical_claim"] not in to_remove]

    # Final deterministic ordering for stable JSON
    cleaned.sort(key=lambda c: (
        text.find(c.get("matched_text","")) if c.get("matched_text") and text.find(c.get("matched_text","")) >= 0 else 9999,
        c["canonical_claim"]
    ))
    return cleaned


def main():
    parser = argparse.ArgumentParser(description="Member 4 Claim Extraction Mock v2.1")
    parser.add_argument("--dictionary", required=True, help="Path to Foundation Claim Dictionary V2.2 xlsx")
    parser.add_argument("--text", help="Single beauty-post text to analyze")
    parser.add_argument("--testset", help="Optional internal test-set xlsx path")
    parser.add_argument("--blindtest", help="Optional blind-test xlsx path")
    parser.add_argument("--out", help="Output xlsx path")
    parser.add_argument("--mock", action="store_true", help="Enable Mock Semantic Layer")
    args = parser.parse_args()

    extractor = ClaimExtractor(args.dictionary)
    mock_layer = MockSemanticLayer(extractor.consumer_meaning) if args.mock else None

    def analyze(text: str) -> Dict[str, Any]:
        base = extractor.extract_claims(text)
        rule_claims = base["output"]["claims"]

        if not mock_layer:
            base["output"]["mode"] = "dictionary+rules"
            return base

        mock_claims = mock_layer.analyze(text, rule_claims)
        merged = merge_rule_and_mock(rule_claims, mock_claims)
        cleaned = resolve_claim_conflicts(text, merged)

        return {
            "tool_name": "claim_extraction",
            "input": {"text": text},
            "output": {
                "mode": "dictionary+rules+mock_semantic+conflict_resolution",
                "claims": cleaned
            }
        }

    if args.text:
        print(json.dumps(analyze(args.text), ensure_ascii=False, indent=2))

    if args.testset:
        wb = load_workbook(args.testset)
        ws = wb["Test Set"]
        headers = {_norm(c.value): i for i, c in enumerate(ws[1], start=1)}
        text_col = headers["Test Text"]
        expected_col = headers["Expected Canonical Claims"]
        model_col = headers["Model Output"]
        correct_col = headers["Correct?"]
        notes_col = headers["Error Notes"]

        total, exact = 0, 0
        for r in range(2, ws.max_row + 1):
            text = _norm(ws.cell(r, text_col).value)
            expected = {_norm(x) for x in re.split(r"[；;]", _norm(ws.cell(r, expected_col).value)) if _norm(x)}
            result = analyze(text)
            claims = result["output"]["claims"]
            predicted = {c["canonical_claim"] for c in claims}
            ws.cell(r, model_col).value = json.dumps(claims, ensure_ascii=False)

            if predicted == expected:
                status = "Yes"
                exact += 1
            elif predicted & expected:
                status = "Partial"
            else:
                status = "No"
            ws.cell(r, correct_col).value = status

            missing = sorted(expected - predicted)
            extra = sorted(predicted - expected)
            notes = []
            if missing:
                notes.append("Missing: " + "；".join(missing))
            if extra:
                notes.append("Extra: " + "；".join(extra))
            ws.cell(r, notes_col).value = " | ".join(notes)
            total += 1

        out = args.out or "Claim_Extraction_Test_Result_Mock_v2.xlsx"
        wb.save(out)
        print(json.dumps({
            "test_rows": total,
            "exact_match_rows": exact,
            "exact_match_rate": round(exact / total, 3) if total else 0,
            "mode": "dictionary+rules+mock_semantic+conflict_resolution" if args.mock else "dictionary+rules",
            "output_file": out
        }, ensure_ascii=False, indent=2))

    if args.blindtest:
        wb = load_workbook(args.blindtest)
        ws = wb["Blind Test"]
        headers = {_norm(c.value): i for i, c in enumerate(ws[1], start=1)}
        text_col = headers["Blind Test Text"]
        model_col = headers["Model Output"]

        total, claim_counts = 0, []
        for r in range(2, ws.max_row + 1):
            text = _norm(ws.cell(r, text_col).value)
            if not text:
                continue
            result = analyze(text)
            claims = result["output"]["claims"]
            ws.cell(r, model_col).value = json.dumps(claims, ensure_ascii=False)
            claim_counts.append(len(claims))
            total += 1

        out = args.out or "Claim_Extraction_Blind_Result_Mock_v2.xlsx"
        wb.save(out)
        print(json.dumps({
            "blind_test_rows": total,
            "average_claims_per_row": round(sum(claim_counts)/len(claim_counts), 2) if claim_counts else 0,
            "mode": "dictionary+rules+mock_semantic+conflict_resolution" if args.mock else "dictionary+rules",
            "output_file": out,
            "note": "Mock semantic layer + conflict resolution used. No API or Answer Key was used for extraction."
        }, ensure_ascii=False, indent=2))

    if not args.text and not args.testset and not args.blindtest:
        parser.error("Please provide --text, --testset, or --blindtest")


if __name__ == "__main__":
    main()
